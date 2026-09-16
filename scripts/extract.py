# -*- coding: utf-8 -*-
"""엑셀(data/한전_전기요금_2026_HEMS정리.xlsx) → data/data.json
현황판이 쓰는 모든 숫자·문구를 뽑는다. 시트 구조(행 위치)가 바뀌면 이 파일도 같이 고쳐야 한다.
실행: python scripts/extract.py
"""
import json, os, sys
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, 'data', '한전_전기요금_2026_HEMS정리.xlsx')
OUT = os.path.join(ROOT, 'data', 'data.json')


def f(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    D = {}

    # 버전·변경이력
    ws = wb['변경이력']
    D['version'] = {'ver': ws['C1'].value, 'date': str(ws['E1'].value)}
    D['history'] = [dict(ver=r[0], date=str(r[1]), cell=r[2], before=r[3], after=r[4], why=r[5], src=r[6], who=r[7], ok=r[8])
                    for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]

    # 주택용 누진 (주택용!A6:H19)
    ws = wb['주택용']
    D['prog'] = [dict(volt=a[0], season=a[1], tier=a[2], lo=f(a[3]), hi=f(a[4]), base=f(a[5]), rate=f(a[6]), src=a[7])
                 for a in ([c.value for c in ws[r]] for r in range(6, 20))]
    D['tou_house'] = dict(base_kw=4310, kw=3, monthly=12930, super=736.2,
                          rates={'여름·겨울': [138.7, 184.7, 220.5], '봄·가을': [125.8, 153.8, 172.4]},
                          eligibility={'육지': '인증 지열·공기열 설비 설치 주거용 고객 (2026.4.1~; 해당 설비에 일반용 별도 적용 중이면 제외)',
                                       '제주': '주거용 전 가구 (2021.9.1~)'})

    # 공통 부과항목 (요금구성_공통!A5:C8)
    ws = wb['요금구성_공통']
    D['common'] = {ws[f'A{r}'].value: f(ws[f'C{r}'].value) for r in range(5, 9)}

    # 일반/산업/교육 단가 (숨김 시트 포함)
    rows = []
    for sh in ['일반용', '산업용', '교육용']:
        ws = wb[sh]
        for r in ws.iter_rows(min_row=5, values_only=True):
            if r[0] and r[1] and r[4] and f(r[3]) is not None:
                d = dict(kind=r[0], volt=r[1], sel=r[2] if r[2] not in (None, '-') else '', band=r[4], base=f(r[3]),
                         s=f(r[5]), sa=f(r[6]), w=f(r[7]), src=r[8] if sh != '산업용' else r[11])
                if sh == '산업용' and f(r[8]) is not None:
                    d['def'] = [f(r[8]), f(r[9]), f(r[10])]
                rows.append(d)
    D['rates'] = rows

    # 시간대 (시간대_육지vs제주!A5:AB14)
    ws = wb['시간대_육지vs제주']
    D['tou'] = [dict(name=a[0], when=a[1], h=[a[i] for i in range(2, 26)], src=a[26], st=a[27])
                for a in ([c.value for c in ws[r]] for r in range(5, 15))]

    # EV (전기차충전!A5:G10)
    ws = wb['전기차충전']
    D['ev'] = [dict(volt=a[0], base=f(a[1]), band=a[2], s=f(a[3]), sa=f(a[4]), w=f(a[5]), src=a[6])
               for a in ([c.value for c in ws[r]] for r in range(5, 11))]

    # 농사·가로등·심야 (참고)
    ws = wb['농사용_가로등_심야']
    D['misc'] = [dict(kind=a[0], sub=a[1], base=a[2], s=f(a[3]), sa=f(a[4]), w=f(a[5]), note=a[6], src=a[7])
                 for a in ([c.value for c in ws[r]] for r in range(5, 13))]

    # 2026 변동 요약 (의사결정요약!A33:F43)
    ws = wb['의사결정요약']
    D['policy'] = [dict(item=a[0], state=a[1], date=a[2], target=a[3], note=a[4], src=a[5])
                   for a in ([c.value for c in ws[r]] for r in range(33, 44))]

    # 캐시백·DR·ESS (ESS_DR_인센티브)
    ws = wb['ESS_DR_인센티브']
    D['cashback'] = [dict(prog=a[0], when=a[1], cond=a[2], pay=a[3], hems=a[4], src=a[5])
                     for a in ([c.value for c in ws[r]] for r in range(13, 16))]
    D['dr'] = [dict(name=ws[f'A{r}'].value, region=ws[f'B{r}'].value, what=ws[f'C{r}'].value, state=ws[f'D{r}'].value, src=ws[f'F{r}'].value) for r in (19, 20)]
    D['ess'] = [dict(period=ws[f'A{r}'].value, energy=ws[f'B{r}'].value, base=ws[f'C{r}'].value, target=ws[f'D{r}'].value, state=ws[f'E{r}'].value, src=ws[f'F{r}'].value) for r in (6, 7, 8)]

    # 출처
    ws = wb['출처']
    D['sources'] = {ws[f'A{r}'].value: dict(name=ws[f'B{r}'].value, org=ws[f'C{r}'].value, date=str(ws[f'D{r}'].value), grade=ws[f'E{r}'].value, url=ws[f'F{r}'].value)
                    for r in range(5, ws.max_row + 1) if ws[f'A{r}'].value}

    json.dump(D, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f"extract ok: {D['version']} rates={len(rows)} history={len(D['history'])} -> {OUT}")


if __name__ == '__main__':
    main()
