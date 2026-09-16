# -*- coding: utf-8 -*-
"""v13: KPX 2026년 6월 수요자원거래시장 현황(2026.9.9 게시, 참여현황은 2026.8월 기준) 반영.
DR_수요반응 시트 갱신 + H절(현황판 키값) 신설 + 변경이력 + 출처 S39.
일회성 스크립트. 실행 후 build.py --all."""
import openpyxl, shutil, os, datetime as dt
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
X = os.path.join(ROOT, 'data', '한전_전기요금_2026_HEMS정리.xlsx')
TODAY = '2026-09-16'
shutil.copy(X, os.path.join(ROOT, 'archive', f'{TODAY}_v12.xlsx'))
wb = openpyxl.load_workbook(X)
F = 'Noto Sans CJK SC'
thin = Side(style='thin', color='FFBFBFBF'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
H_FILL = PatternFill('solid', fgColor='FF1F4E78'); SEC = PatternFill('solid', fgColor='FFD9E1F2'); KEY = PatternFill('solid', fgColor='FFF2F2F2'); EST = PatternFill('solid', fgColor='FFFFF2CC')

ws = wb['DR_수요반응']
log = []  # (cell, before, after, why, src)


def setv(addr, val, why, src='S39'):
    c = ws[addr]; log.append((f'DR_수요반응!{addr}', c.value, val, why, src)); c.value = val


# 머리말 근거
ws['A2'].value = ws['A2'].value.replace('「2025년 12월 수요자원거래시장 현황 및 운영정보」(2026.3 발간, S33)', '「2026년 6월 수요자원거래시장 현황 및 운영정보」(2026.9.9 게시, 참여현황 2026.8월 기준, S39; 이전 2025.12월분 S33)')
# A절
setv('F6', '국민DR 100,789고객 · 육지 플러스DR 42,651고객 · 제주 플러스DR 1,310고객 (2026.8)', '참여현황 갱신')
setv('F7', '28개 사업자 · 81자원 · 4,033고객 · 4,326MW / 휴일DR 9자원 · 132고객 · 975MW (2026.8)', '참여현황 갱신')
ws['G6'].value = 'S35 제12.1.3조 · S39 p.5'; ws['G7'].value = 'S35 제12.2.2조 · S39 p.3'
# B절
setv('A23', '단가 (KPX 공표)', '연도 갱신')
setv('B23', '2026년 RBP 육지 65,760원/kW (1월 65,871), 제주 74,771원/kW — 2025년(67,277/76,562) 대비 약 2% 인하. 실적 평균 2025 ≈ 1,625원/kWh, 2026 상반기 ≈ 1,552원/kWh (10.59억 ÷ 682,625kWh)', '2026 단가 반영')
setv('C23', '육지 50원/kWh. 단, 2026.4.16~ 봄·가을철 공휴일·토요일 및 설·추석 연휴는 별표26 산식으로 61.35원/kWh (규칙 26-긴급4차). 제주 2026.1~5월 96원 → 6월 76원', '휴일 단가 61.35 신설')
setv('D23', '고정기본정산금 2026.6 1,952원/kW-월 (2026.1~6: 744~2,650)', '2026 단가')
setv('F23', '2,446원/kW-월 (2026.1~6: 907~3,286)', '2026 단가')
setv('G23', '224~972원/kW-월 (2026.2~5)', '2026 단가')
setv('A24', '2026.8 참여 현황', '기준월 갱신')
setv('B24', '12개 사업자 · 218자원 · 100,789고객 (2026.2 89,973 → 8월 10만 돌파)', '참여현황')
setv('C24', '육지 10사업자·172자원·42,651고객·2,110MW / 제주 7사업자·25자원·1,310고객', '참여현황')
setv('D24', '수도권 609MW·553고객, 비수도권 3,247MW·1,265고객', '참여현황')
setv('E24', '수도권 196MW·828고객, 비수도권 253MW·1,318고객', '참여현황')
setv('F24', '22MW·69고객', '참여현황')
setv('G24', '표준 수도권 142MW·13고객, 비수도권 833MW·119고객 (휴일 합계 975MW·132고객)', '참여현황')
setv('H24', '5단계 체계(2025.4.10~): 1단계 639MW·58고객, 2단계 314MW, 3단계 418MW·32고객, 4단계 255MW, 5단계 333MW', '참여현황')
setv('A25', '실적 (2025 / 2026 상반기)', '연도 갱신')
setv('B25', '2025: 45시간 · 798,187kWh · 12.97억원 / 2026 1~6월: 38시간 · 682,625kWh · 10.59억원 (누적 2020~26.6: 344회, 1,661MWh, 26.7억)', '2026 실적')
setv('C25', '육지 2025: 16일 63시간 30,725MWh 13.1억 → 2026 3~6월: 26일 104시간 47,561MWh 21.6억 (전년 대비 1.7배). 제주: 2023 이후 표 미수록', '2026 실적')
setv('D25', '자발적DR 2026 1~6월 낙찰 6,963~17,899MWh/월, 감축 10,735~25,612MWh/월. 미세먼지DR은 규칙 26-정기1차로 폐지', '2026 실적·미세먼지DR 폐지')
setv('A28', ws['A28'].value.replace('"2025년 단가" 행의 국민DR 평균 1,625원/kWh', '"단가" 행의 국민DR 평균(2025 1,625 / 2026 상반기 1,552원/kWh)'), '주석 갱신')
# C절
setv('B41', '2025년: 1,297,300천원 ÷ 798,187kWh ≈ 1,625원/kWh. 2026년 1~6월: 1,059,100천원 ÷ 682,625kWh ≈ 1,552원/kWh (RBP 인하 반영). 2026.6.29 19시 발령: 24,290kWh → 37,146천원(1,529원/kWh)', '2026 실적 역산')
setv('B43', '2025: 45시간, 798,187kWh, 12.97억원. 2026 1~6월: 38시간, 682,625kWh, 10.59억원. 참여고객 2025.12 85,011 → 2026.2 89,973 → 2026.8 100,789 (12개 사업자, 218자원)', '2026 실적·참여')
ws['C43'].value = 'S39 p.5·6·17·18'
# D절
setv('A52', '정산 단가 (2025 → 2026)', '연도 갱신')
setv('B52', '육지 50원/kWh (2025 연중, 2026 평일). 2026.4.16부터 봄·가을철 공휴일·토요일 및 설·추석 연휴는 별표26 산식 61.35원/kWh(규칙 26-긴급4차). 제주 2025 95~97 → 2026.1~5월 96 → 6월 76원/kWh. 증대량 × 단가', '휴일 61.35 신설')
ws['C52'].value = 'S39 p.10'
setv('A53', '실적 (2025 → 2026 상반기)', '연도 갱신')
setv('B53', '육지 2025: 3~6월 16일 63시간 30,725MWh 13.1억. 2026: 3월 4일13h 7,085MWh, 4월 7일33h 19,409MWh, 5월 10일42h 15,481MWh, 6월 5일16h 5,586MWh = 26일 104시간 47,561MWh 21.6억 (2026.6 발령 예: 6/1 11~15시, 6/13 11~14시, 6/14 12~14시, 6/28 11~13시 — 주말 낮 집중). 참여고객 육지 42,651 / 제주 1,310 (2026.8), 육지 증대가능용량 2,110MW', '2026 실적')
ws['C53'].value = 'S39 p.5·17·19'
# F절
setv('B67', '수도권/비수도권 자원. RBP 65,760원/kW (2026, 2025는 67,277)', '2026 단가')
setv('C67', '제주권 자원. RBP 74,771원/kW (2026) — 육지보다 14% 높음', '2026 단가')
ws['D67'].value = 'S39 p.10'
setv('B68', '2024.3.28 신설. 증대필요량 = 태양광이용률·수요구간 기준(KPX 매월 공지). 평일 50원/kWh, 봄·가을 휴일·명절 61.35원/kWh(2026.4.16~). 42,651고객. 2026 상반기 104시간 발령', '2026 반영')
setv('C68', '2021.1 신설, 2022.5 시장 분리. 출력제어량 기준 낙찰, 실시간(당일) 시장 병행. 2026 96 → 6월 76원/kWh. 1,310고객', '2026 반영')
ws['D68'].value = 'S35 제12.7.3.6~3.8조, S39'
setv('B70', '국민DR + 플러스DR 둘 다 가능. 플러스DR은 2026년 봄·가을 주말 낮에 집중 발령(104시간)되고 휴일 단가 61.35원이라 EV·ESS 있는 가정은 주말 낮 충전 스케줄과 결합 가치 있음', '2026 반영')
# G절
setv('B75', '가정이 1시간에 1kWh 줄이면 KPX 정산 ≈ 1,550~1,600원(2025 1,625 / 2026 상반기 1,552). 연 40~45시간 전부 참여·매회 1kWh 감축 시 KPX 정산 ≈ 6~7만원, 사업자 배분 후 가정 수령은 그 일부', '2026 반영')
setv('B76', 'EV 30kWh 충전을 잉여 시간으로 옮기면 육지 평일 50원 × 30 = 1,500원/회, 봄·가을 휴일 61.35 × 30 = 1,841원/회, 제주 76~96원 × 30 = 2,280~2,880원/회 (KPX 정산 기준). 2026 상반기 육지 발령 104시간이므로 연 수만 원 규모. 요금 측면 이득이 DR 정산보다 큼', '2026 반영')
setv('B78', '2026 하반기 RBP·MRT(KPX 월간 현황 보고서로 매월 갱신), 사업자별 고객 배분율·보상 방식, 제주 플러스DR 2024~2026 실적(보고서 표 미수록), 61.35원 적용 산식(별표26) 원문', '갱신')

# H절: 현황판 키값 (기계가 읽는 표)
r = ws.max_row + 2
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
c = ws.cell(r, 1, 'H. 현황판 키값 (scripts/extract.py가 읽어 현황판에 표시 — 키 이름 변경 금지, 값만 갱신)')
c.font = Font(name=F, size=10, bold=True); c.fill = SEC; c.alignment = Alignment(vertical='center')
for i in range(1, 11): ws.cell(r, i).border = B
r += 1
for i, t in enumerate(['키', '값', '단위·설명', '근거', '상태'], 1):
    c = ws.cell(r, i, t); c.font = Font(name=F, size=10, bold=True, color='FFFFFFFF'); c.fill = H_FILL; c.alignment = Alignment(horizontal='center', vertical='center'); c.border = B
r += 1
KV = [
    ('기준월', '2026.8', 'KPX 보고서의 참여현황 기준월', 'S39 p.5', '확정'),
    ('국민DR_참여고객', 100789, '명', 'S39 p.5', '확정'),
    ('국민DR_사업자', 12, '개', 'S39 p.5', '확정'),
    ('국민DR_자원', 218, '개', 'S39 p.5', '확정'),
    ('국민DR_RBP_육지', 65760, '원/kW (2026)', 'S39 p.10', '확정'),
    ('국민DR_RBP_제주', 74771, '원/kW (2026)', 'S39 p.10', '확정'),
    ('국민DR_실적평균단가', 1552, '원/kWh (2026 상반기 역산)', 'S39 p.18', '추정(계산식 병기)'),
    ('국민DR_실적문구', '2026 1~6월 38시간 · 683MWh (2025 45시간 · 798MWh)', '현황판 표시용', 'S39 p.18', '확정'),
    ('플러스DR_단가_육지_평일', 50, '원/kWh', 'S39 p.10', '확정'),
    ('플러스DR_단가_육지_휴일', 61.35, '원/kWh (2026.4.16~ 봄·가을 공휴일·토요일·설·추석)', 'S39 p.10', '확정'),
    ('플러스DR_단가_제주', 76, '원/kWh (2026.6; 1~5월 96)', 'S39 p.10', '확정'),
    ('플러스DR_고객_육지', 42651, '명', 'S39 p.5', '확정'),
    ('플러스DR_고객_제주', 1310, '명', 'S39 p.5', '확정'),
    ('플러스DR_용량_육지', 2110, 'MW (증대가능용량)', 'S39 p.5', '확정'),
    ('플러스DR_실적문구_육지', '2026 3~6월 26일 104시간 47,561MWh (2025 63시간)', '현황판 표시용', 'S39 p.19', '확정'),
    ('플러스DR_실적문구_제주', '2023 이후 보고서 표 미수록', '현황판 표시용', 'S39 p.20', '확인필요'),
    ('사업장DR_요약', '28사업자 · 81자원 · 4,033고객 · 4,326MW / 휴일DR 975MW·132고객', '현황판 표시용', 'S39 p.3', '확정'),
]
for k in KV:
    for i, v in enumerate(k, 1):
        c = ws.cell(r, i, v); c.font = Font(name=F, size=10, bold=(i == 1)); c.border = B; c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        if i == 1: c.fill = KEY
        if i == 5 and '추정' in str(v): ws.cell(r, 2).fill = EST
        if i == 5 and '확인필요' in str(v): ws.cell(r, 2).fill = PatternFill('solid', fgColor='FFFCE4D6')
    r += 1
log.append((f'DR_수요반응!A{r-len(KV)-2}:E{r-1}', '—', 'H절 현황판 키값 17행 신설', '현황판이 엑셀에서 DR 숫자를 읽도록', 'S39'))

# 출처 S39
O = wb['출처']; rr = O.max_row + 1
vals = ['S39', '2026년 6월 수요자원거래시장 현황 및 운영정보 (2026.9.9 게시; 참여현황 2026.8월, 단가 2026.1~6월, 실적 2026.6월·누적)', '전력거래소', '2026-09-09', 'A', 'https://kpx.or.kr/board.es?mid=a10102000000&bid=0088&act=view&list_no=78099 (evidence/KPX_2026-06_수요자원시장현황.pdf)']
for i, col in enumerate('ABCDEF'):
    s = O[f'{col}36']; d = O[f'{col}{rr}']
    d.font = s.font.copy(); d.fill = s.fill.copy(); d.alignment = s.alignment.copy(); d.border = s.border.copy(); d.value = vals[i]
O.row_dimensions[rr].height = 30

# 변경이력
H = wb['변경이력']; H['C1'] = 'v13'; H['E1'] = TODAY
hr = H.max_row + 1
for cell, before, after, why, src in log:
    row = ['v13', TODAY, cell, (str(before)[:120] if before else '—'), str(after)[:160], why, src, 'Claude', '사용자']
    for i, v in enumerate(row, 1):
        c = H.cell(hr, i, v); c.font = Font(name=F, size=10, bold=(i == 1)); c.border = B; c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        if i == 1: c.fill = KEY
    H.row_dimensions[hr].height = 40; hr += 1
S = wb['의사결정요약']; S['A2'] = S['A2'].value.replace('기준일 2026-09-15', f'기준일 {TODAY}')
wb.calculation.fullCalcOnLoad = True
wb.save(X)
print('v13 saved; changes:', len(log))
